import os

from PyQt6 import uic
from PyQt6.QtWidgets import (QApplication, QMainWindow, QPushButton, QToolButton, QFileDialog, QLineEdit,
                             QCheckBox, QLabel, QMessageBox)
import sys
import numpy as np
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import psycopg2
import resources_rc
import PostgreSQL
from CatBoost_Classifier import ThesisClassifier
from Graphs import ThesisVisualizer



theses_files = []
categories_file = None
authors_file = None
MODEL_PATH = "thesis_classifier_catboost.pkl"

class App(QMainWindow):
    def __init__(self):
        super().__init__()
        uic.loadUi('Program Interface.ui', self)

        self.loadThesesButton.clicked.connect(self.load_theses)
        self.loadCategButton.clicked.connect(self.load_categories)
        self.loadAuthorsButton.clicked.connect(self.load_authors)

        self.clearThesesButton.clicked.connect(self.clear_theses)
        self.clearCategButton.clicked.connect(self.clear_categories)
        self.clearAuthorsButton.clicked.connect(self.clear_authors)

        self.viewThesesButton.clicked.connect(self.view_theses)
        self.viewCategButton.clicked.connect(self.view_categories)
        self.viewAuthorsButton.clicked.connect(self.view_authors)

        self.HelpButton.clicked.connect(self.show_help)
        self.StartButton.clicked.connect(self.classification_start)

#region Загрузка файлов

    def load_theses(self):
        files = QFileDialog.getOpenFileNames(
            parent=self,
            caption='Выберите файл(ы)',
            directory=os.getcwd(),
            filter='Word Files (*.docx);; Text Files (*.txt);; All Files (*)',
            initialFilter='Text File (*.txt)'
        )
        notLoaded = 0
        if files[0]:
            for file in files[0]:
                if (file.endswith('.txt') or file.endswith('.docx')) and file not in theses_files:
                    theses_files.append(file)
                elif not (file.endswith('.txt') or file.endswith('.docx')):
                    notLoaded = 2
        else:
            notLoaded = 1

        if notLoaded == 0:
            self.debugThesesLabel.setText("Файл(ы) успешно загружены")
        elif notLoaded == 1 and len(theses_files) == 0:
            self.debugThesesLabel.setText("Загрузите как минимум один файл!")
        else:
            self.debugThesesLabel.setText("Некорректный формат файла(ов)!")

        pass

    def load_categories(self):
        global categories_file
        file = QFileDialog.getOpenFileName(
            parent=self,
            caption='Выберите файл',
            directory=os.getcwd(),
            filter='Word Files (*.docx);; Text Files (*.txt);; All Files (*)',
            initialFilter='Text File (*.txt)'
        )
        notLoaded = 0
        if file[0]:
            if categories_file is None:
                if file[0].endswith('.txt') or file[0].endswith('.docx'):
                    categories_file = file[0]
                elif not (file[0].endswith('.txt') or file[0].endswith('.docx')):
                    notLoaded = 2
            else:
                notLoaded = 3
        else:
            notLoaded = 1

        if notLoaded == 0:
            self.debugCategLabel.setText("Файл успешно загружен")
        elif notLoaded == 1:
            self.debugCategLabel.setText("Загрузите один файл!")
        elif notLoaded == 2:
            self.debugCategLabel.setText("Некорректный формат файла!")
        else:
            self.debugCategLabel.setText("Такой же или другой файл уже загружен!")

        pass

    def load_authors(self):
        global authors_file
        file = QFileDialog.getOpenFileName(
            parent=self,
            caption='Выберите файл',
            directory=os.getcwd(),
            filter='Word Files (*.docx);; Text Files (*.txt);; All Files (*)',
            initialFilter='Text File (*.txt)'
        )
        notLoaded = 0
        if file[0]:
            if authors_file is None:
                if file[0].endswith('.txt') or file[0].endswith('.docx'):
                    authors_file = file[0]
                else:
                    notLoaded = 3
            else:
                notLoaded = 2
        else:
            notLoaded = 1

        if notLoaded == 0:
            self.debugAuthorsLabel.setText("Файл успешно загружен")
        elif notLoaded == 1:
            self.debugAuthorsLabel.setText("Это необязательное поле")
        elif notLoaded == 2:
            self.debugAuthorsLabel.setText("Такой же или другой файл уже загружен!")
        else:
            self.debugAuthorsLabel.setText("Некорректный формат файла!")

        pass

#endregion


    def clear_theses(self):
        global theses_files
        theses_files = []
        self.debugThesesLabel.setText("Должен быть загружен хотя бы один файл!")
        pass

    def clear_categories(self):
        global categories_file
        categories_file = None
        self.debugCategLabel.setText("Должен быть загружен один файл!")
        pass

    def clear_authors(self):
        global authors_file
        authors_file = None
        self.debugAuthorsLabel.setText("Это необязательное поле")
        pass


    # Показать список загруженных файлов тезисов
    def view_theses(self):
        if not theses_files:
            QMessageBox.information(self, "Загруженные тезисы", "Файлы тезисов не загружены.")
            return

        files_list = "\n".join([f"{i + 1}. {f}" for i, f in enumerate(theses_files)])
        QMessageBox.information(
            self,
            "Загруженные тезисы",
            f"Загружено файлов: {len(theses_files)}\n\n{files_list}"
        )

    # Показать загруженный файл категорий
    def view_categories(self):
        if categories_file is None:
            QMessageBox.information(self, "Загруженные категории", "Файл категорий не загружен.")
            return

        QMessageBox.information(
            self,
            "Загруженные категории",
            f"Загруженный файл:\n{categories_file}"
        )

    # Показать загруженный файл авторов с полным путём
    def view_authors(self):
        if authors_file is None:
            QMessageBox.information(self, "Загруженные авторы", "Файл авторов не загружен.")
            return

        QMessageBox.information(
            self,
            "Загруженные авторы",
            f"Загруженный файл:\n{authors_file}"
        )


    # Показать окно помощи с описанием функций программы
    def show_help(self):
        help_text = """
        <h2>Справка по программе классификации тезисов</h2>

        <h3>Основные функции:</h3>

        <h4>1. Загрузка данных</h4>
        <ul>
            <li><b>Загрузить тезисы</b> - загрузка файлов с тезисами (txt/docx). Можно выбрать несколько файлов</li>
            <li><b>Загрузить категории</b> - загрузка файла со списком категорий (один файл)</li>
            <li><b>Загрузить авторов</b> - загрузка файла со списком авторов (опционально)</li>
        </ul>

        <h4>2. Очистка данных</h4>
        <ul>
            <li><b>Очистить тезисы</b> - удаляет все загруженные тезисы</li>
            <li><b>Очистить категории</b> - удаляет загруженные категории</li>
            <li><b>Очистить авторов</b> - удаляет загруженных авторов</li>
        </ul>

        <h4>3. Настройки классификации</h4>
        <ul>
            <li><b>Сохранить результат в файл Word</b> - экспорт результатов в документ Word</li>
            <li><b>Сохранить результат в файл TXT</b> - экспорт результатов в текстовый файл</li>
            <li><b>Сохранить данные распределения</b> - сохранение детальной статистики классификации</li>
            <li><b>Вывести график распределения</b> - визуализация распределения тезисов по категориям</li>
            <li><b>Выполнить балансировку категорий</b> - равномерное распределение тезисов по категориям</li>
        </ul>

        <h4>4. Настройки подключения к БД</h4>
        <ul>
            <li>Укажите параметры подключения к PostgreSQL:</li>
            <li>Хост, Порт, База данных, Пользователь, Пароль</li>
        </ul>

        <h3>Порядок работы:</h3>
        <ol>
            <li>Укажите параметры подключения к базе данных</li>
            <li>Загрузите файл с категориями</li>
            <li>Загрузите файлы с тезисами</li>
            <li>При необходимости загрузите файл с авторами</li>
            <li>Выберите необходимые опции</li>
                <li>Нажмите "Запуск"</li>
        </ol>

        <h3>Примечания:</h3>
        <ul>
            <li>Поддерживаются файлы форматов .txt и .docx</li>
            <li>Для работы программы необходим установленный PostgreSQL</li>
            <li>Модель классификации должна находиться в файле thesis_classifier_catboost.pkl</li>
        </ul>
        """

        QMessageBox.about(self, "Справка", help_text)


    def classification_start(self):
        DB_CONFIG = {
            'host': self.hostLine.text(),
            'port': self.portLine.text(),
            'database': self.databaseLine.text(),
            'user': self.userLine.text(),
            'password': self.passwordLine.text()
        }

        PostgreSQL.create_tables(DB_CONFIG)

        # Загрузка авторов из файла
        PostgreSQL.text_upload_from_file(1, [authors_file], DB_CONFIG)
        # Загрузка категорий из файла
        PostgreSQL.text_upload_from_file(2, [categories_file], DB_CONFIG)
        # Загрузка тезисов из файлов
        PostgreSQL.text_upload_from_file(3, theses_files, DB_CONFIG)

        # Предобработка текстов
        processed = PostgreSQL.process_texts_from_db(DB_CONFIG, True)
        processed = PostgreSQL.process_texts_from_db(DB_CONFIG, False)

        # Загрузка классификатора
        print("\nЗагрузка модели...")
        classifier = ThesisClassifier(MODEL_PATH, DB_CONFIG)

        # Показ статистики до классификации
        print("\nСтатистика до классификации:")
        stats_before = classifier.get_classification_stats()
        for key, value in stats_before.items():
            if isinstance(value, float):
                print(f"  {key}: {value:.2f}")
            else:
                print(f"  {key}: {value}")

        classifier.update_thesis_predictions(threshold=0.0)

        # Перебалансировка распределения
        if self.doBalanceCheckBox.isChecked():
            target_ratio = 0.6
            min_confidence = 0.0

            classifier.rebalance_theses_distribution(
                target_balance_ratio=target_ratio,
                min_confidence_threshold=min_confidence,
                max_iterations=10)

        # Показываем итоговую статистику
        print("\nИТОГОВАЯ СТАТИСТИКА")
        stats_after = classifier.get_classification_stats()
        for key, value in stats_after.items():
            if isinstance(value, float):
                print(f"  {key}: {value:.2f}")
            else:
                print(f"  {key}: {value}")


        if self.saveWordCheckBox.isChecked():
            try:
                from datetime import datetime

                # Подключаемся к БД
                conn = psycopg2.connect(**DB_CONFIG)
                cursor = conn.cursor()

                # Получаем все категории с тезисами
                cursor.execute("""
                            SELECT 
                                c.id,
                                c.raw_text,
                                COUNT(t.id) as thesis_count
                            FROM categories c
                            INNER JOIN theses t ON t.chosen_category_id = c.id
                            WHERE c.cleaned_text IS NOT NULL AND c.cleaned_text != ''
                            GROUP BY c.id, c.raw_text
                            ORDER BY c.id ASC
                        """)

                categories = cursor.fetchall()

                if not categories:
                    QMessageBox.warning(self, "Предупреждение", "Нет категорий с назначенными тезисами!")
                    return

                # Создаем Word документ
                doc = Document()

                # Настройка стилей
                style = doc.styles['Normal']
                font = style.font
                font.name = 'Times New Roman'
                font.size = Pt(12)

                # Заголовок документа
                title = doc.add_heading('Классификация тезисов по категориям', level=0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER

                # Добавляем информацию о дате
                date_paragraph = doc.add_paragraph()
                date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                date_run = date_paragraph.add_run(f'Дата создания: {datetime.now().strftime("%d.%m.%Y %H:%M")}')
                date_run.font.size = Pt(10)
                date_run.font.color.rgb = RGBColor(128, 128, 128)

                doc.add_paragraph()  # Пустая строка

                # Счетчик тезисов
                total_theses = 0

                # Проходим по всем категориям
                for cat_id, cat_name, thesis_count in categories:
                    # Сначала получаем уникальные тезисы
                    cursor.execute("""
                        SELECT DISTINCT
                            t.id,
                            t.raw_text,
                            t.cleaned_text
                        FROM theses t
                        WHERE t.chosen_category_id = %s
                        ORDER BY t.id
                    """, (cat_id,))

                    theses = cursor.fetchall()

                    if not theses:
                        continue

                    # Добавляем заголовок категории
                    category_heading = doc.add_heading(f'{cat_name}', level=1)

                    # Информация о количестве тезисов
                    info_paragraph = doc.add_paragraph()
                    info_run = info_paragraph.add_run(f'Количество тезисов: {len(theses)}')
                    info_run.font.size = Pt(10)
                    info_run.font.color.rgb = RGBColor(100, 100, 100)
                    info_run.italic = True

                    # Добавляем разделительную линию
                    doc.add_paragraph('_' * 50)

                    # Добавляем каждый тезис
                    for thesis_id, raw_text, cleaned_text in theses:
                        # Номер тезиса
                        thesis_paragraph = doc.add_paragraph()
                        thesis_number = thesis_paragraph.add_run(f'Тезис #{thesis_id}')
                        thesis_number.bold = True
                        thesis_number.font.size = Pt(11)

                        # Текст тезиса
                        text_paragraph = doc.add_paragraph()
                        text_run = text_paragraph.add_run(raw_text if raw_text else 'Текст отсутствует')
                        text_run.font.size = Pt(11)

                        # Отступ между тезисами
                        doc.add_paragraph()

                        total_theses += 1

                    # Добавляем разделитель между категориями
                    doc.add_paragraph('=' * 50)
                    doc.add_paragraph()

                # Добавляем итоговую статистику
                doc.add_page_break()
                stats_heading = doc.add_heading('Статистика классификации', level=1)

                # Общая статистика
                stats_paragraph = doc.add_paragraph()
                stats_paragraph.add_run('Общая статистика:').bold = True
                stats_paragraph.add_run(f'\nВсего тезисов: {total_theses}')
                stats_paragraph.add_run(f'\nВсего категорий: {len(categories)}')
                stats_paragraph.add_run(
                    f'\nСреднее количество тезисов в категории: {total_theses / len(categories):.1f}')

                # Статистика по категориям
                doc.add_heading('Распределение по категориям', level=2)

                for cat_id, cat_name, thesis_count in categories:
                    percentage = (thesis_count / total_theses * 100) if total_theses > 0 else 0
                    stats_paragraph = doc.add_paragraph()
                    stats_paragraph.add_run(f'• {cat_name}: ').bold = True
                    stats_paragraph.add_run(f'{thesis_count} тезисов ({percentage:.1f}%)')
                    stats_paragraph.style.font.size = Pt(10)

                # Сохраняем файл
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"classified_theses_{timestamp}.docx"
                doc.save(filename)

                # Закрываем соединение с БД
                cursor.close()
                conn.close()

                QMessageBox.information(self, "Успех", f"Тезисы сохранены в файл:\n{filename}")

            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Ошибка при сохранении Word файла: {str(e)}")
                print(f"Ошибка при сохранении Word файла: {e}")
                import traceback
                traceback.print_exc()



        if self.saveTXTCheckBox.isChecked():
            try:
                from datetime import datetime

                # Подключаемся к БД
                conn = psycopg2.connect(**DB_CONFIG)
                cursor = conn.cursor()

                # Получаем все категории с назначенными тезисами
                cursor.execute("""
                    SELECT 
                        c.id,
                        c.raw_text,
                        COUNT(t.id) as thesis_count
                    FROM categories c
                    INNER JOIN theses t ON t.chosen_category_id = c.id
                    WHERE c.cleaned_text IS NOT NULL AND c.cleaned_text != ''
                    GROUP BY c.id, c.raw_text
                    ORDER BY c.id ASC
                """)

                categories = cursor.fetchall()

                if not categories:
                    QMessageBox.warning(self, "Предупреждение",
                                        "Нет категорий с назначенными тезисами!")
                    cursor.close()
                    conn.close()
                    return

                # Имя файла с меткой времени
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"classified_theses_{timestamp}.txt"

                total_theses = 0

                # Открываем файл для записи
                with open(filename, 'w', encoding='utf-8') as f:
                    # Заголовок
                    f.write("=" * 80 + "\n")
                    f.write("КЛАССИФИКАЦИЯ ТЕЗИСОВ ПО КАТЕГОРИЯМ\n")
                    f.write(f"Дата создания: {datetime.now().strftime('%d.%m.%Y %H:%M')}\n")
                    f.write("=" * 80 + "\n\n")

                    # Категории и тезисы
                    for cat_id, cat_name, thesis_count in categories:
                        cursor.execute("""
                            SELECT DISTINCT
                                t.id,
                                t.raw_text,
                                t.cleaned_text
                            FROM theses t
                            WHERE t.chosen_category_id = %s
                            ORDER BY t.id
                        """, (cat_id,))

                        theses = cursor.fetchall()
                        if not theses:
                            continue

                        # Заголовок категории
                        f.write("\n" + "=" * 80 + "\n")
                        f.write(f"КАТЕГОРИЯ: {cat_name}\n")
                        f.write(f"Количество тезисов: {len(theses)}\n")
                        f.write("=" * 80 + "\n\n")

                        # Тезисы
                        for thesis_id, raw_text, cleaned_text in theses:
                            f.write(f"--- Тезис #{thesis_id} ---\n")
                            f.write((raw_text if raw_text else "Текст отсутствует") + "\n\n")
                            total_theses += 1

                    # Итоговая статистика
                    f.write("\n" + "#" * 80 + "\n")
                    f.write("ИТОГОВАЯ СТАТИСТИКА\n")
                    f.write("#" * 80 + "\n")
                    f.write(f"Всего тезисов: {total_theses}\n")
                    f.write(f"Всего категорий: {len(categories)}\n")
                    if categories:
                        f.write(f"Среднее количество тезисов в категории: "
                                f"{total_theses / len(categories):.1f}\n")

                    f.write("\nРаспределение по категориям:\n")
                    f.write("-" * 80 + "\n")
                    for cat_id, cat_name, thesis_count in categories:
                        percentage = (thesis_count / total_theses * 100) if total_theses > 0 else 0
                        f.write(f"  {cat_name}: {thesis_count} тезисов ({percentage:.1f}%)\n")

                cursor.close()
                conn.close()

                QMessageBox.information(self, "Успех",
                                        f"Тезисы сохранены в файл:\n{filename}")
                print(f"TXT файл сохранён: {filename}")

            except Exception as e:
                QMessageBox.critical(self, "Ошибка",
                                     f"Ошибка при сохранении TXT файла: {str(e)}")
                print(f"Ошибка при сохранении TXT файла: {e}")
                import traceback
                traceback.print_exc()



        if self.saveDataCheckBox.isChecked():
            try:
                # Создаем экземпляр визуализатора для получения детальной статистики
                visualizer = ThesisVisualizer(DB_CONFIG)

                # Получаем распределение
                distribution = visualizer.get_category_distribution()
                initial_distribution = visualizer.get_initial_distribution()

                # Создаем имя файла с временной меткой
                from datetime import datetime
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"classification_statistics_{timestamp}.txt"

                # Открываем файл для записи
                with open(filename, 'w', encoding='utf-8') as f:
                    # Заголовок
                    f.write("=" * 80 + "\n")
                    f.write("СТАТИСТИКА КЛАССИФИКАЦИИ ТЕЗИСОВ\n")
                    f.write(f"Дата и время: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                    f.write("=" * 80 + "\n\n")

                    # Общая статистика
                    f.write("ОБЩАЯ СТАТИСТИКА\n")
                    f.write("-" * 40 + "\n")
                    for key, value in stats_after.items():
                        if isinstance(value, float):
                            f.write(f"{key}: {value:.2f}\n")
                        else:
                            f.write(f"{key}: {value}\n")

                    f.write("\n" + "=" * 80 + "\n")
                    f.write("РАСПРЕДЕЛЕНИЕ ТЕЗИСОВ ПО КАТЕГОРИЯМ\n")
                    f.write("=" * 80 + "\n\n")

                    # Заголовки таблицы
                    f.write(f"{'ID':<5} {'Категория':<50} {'Кол-во':<8} {'Ср. вероятность':<15}\n")
                    f.write("-" * 80 + "\n")

                    # Данные по категориям
                    for cat_id, cat_text, count, avg_prob in distribution:
                        # Обрезаем длинный текст категории
                        short_text = cat_text[:47] + "..." if len(cat_text) > 47 else cat_text
                        f.write(f"{cat_id:<5} {short_text:<50} {count:<8} {avg_prob:<15.4f}\n")

                    # Итоговая статистика распределения
                    counts = [count for _, _, count, _ in distribution]
                    probs = [prob for _, _, _, prob in distribution if prob > 0]

                    f.write("\n" + "=" * 80 + "\n")
                    f.write("СТАТИСТИКА РАСПРЕДЕЛЕНИЯ\n")
                    f.write("=" * 80 + "\n")
                    f.write(f"Всего категорий: {len(distribution)}\n")
                    f.write(f"Категорий с тезисами: {len([c for c in counts if c > 0])}\n")
                    f.write(f"Всего тезисов: {sum(counts)}\n")
                    f.write(f"Среднее количество тезисов в категории: {np.mean(counts):.2f}\n")
                    f.write(f"Медиана: {np.median(counts):.2f}\n")
                    f.write(f"Стандартное отклонение: {np.std(counts):.2f}\n")
                    f.write(f"Минимум: {min(counts)}\n")
                    f.write(f"Максимум: {max(counts)}\n")

                    if probs:
                        f.write(f"\nСредняя вероятность: {np.mean(probs):.4f}\n")
                        f.write(f"Медианная вероятность: {np.median(probs):.4f}\n")

                    # Сравнение с начальным распределением, если доступно
                    if initial_distribution:
                        f.write("\n" + "=" * 80 + "\n")
                        f.write("СРАВНЕНИЕ С НАЧАЛЬНЫМ РАСПРЕДЕЛЕНИЕМ\n")
                        f.write("=" * 80 + "\n\n")
                        f.write(f"{'ID':<5} {'Категория':<40} {'Было':<8} {'Стало':<8} {'Изменение':<10}\n")
                        f.write("-" * 80 + "\n")

                        init_counts = {cat_id: count for cat_id, _, count, _ in initial_distribution}

                        for cat_id, cat_text, curr_count, _ in distribution:
                            init_count = init_counts.get(cat_id, 0)
                            change = curr_count - init_count
                            short_text = cat_text[:37] + "..." if len(cat_text) > 37 else cat_text
                            f.write(f"{cat_id:<5} {short_text:<40} {init_count:<8} {curr_count:<8} {change:+<10}\n")

                    # Топ-10 категорий
                    f.write("\n" + "=" * 80 + "\n")
                    f.write("ТОП-10 КАТЕГОРИЙ ПО КОЛИЧЕСТВУ ТЕЗИСОВ\n")
                    f.write("=" * 80 + "\n")

                    sorted_dist = sorted(distribution, key=lambda x: x[2], reverse=True)
                    for i, (cat_id, cat_text, count, avg_prob) in enumerate(sorted_dist[:10], 1):
                        short_text = cat_text[:50] + "..." if len(cat_text) > 50 else cat_text
                        f.write(f"{i:2d}. ID:{cat_id:<5} {short_text:<50} {count:>5} тез.\n")

                visualizer.close()

                QMessageBox.information(self, "Успех", f"Статистика сохранена в файл:\n{filename}")
                print(f"Статистика сохранена в файл: {filename}")

            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Ошибка при сохранении статистики: {str(e)}")
                print(f"Ошибка при сохранении статистики: {e}")


        if self.doGraphCheckBox.isChecked():
            try:
                # Создаем экземпляр визуализатора с конфигурацией из интерфейса
                visualizer = ThesisVisualizer(DB_CONFIG)

                # Выводим статистику
                visualizer.print_statistics()

                # Проверяем наличие начального распределения
                initial_dist = visualizer.get_initial_distribution()

                if initial_dist:
                    # Если есть данные о начальном распределении, строим сравнительный график
                    visualizer.plot_comparison_charts(
                        top_n=None,
                        save_path='category_distribution_comparison.png',
                        show_values=True,
                        figsize=(18, 8)
                    )
                else:
                    # Иначе строим только текущий график
                    visualizer.plot_bar_chart(
                        top_n=None,
                        save_path='category_distribution_current.png',
                        show_values=True,
                        figsize=(12, 8)
                    )

                # Закрываем соединение
                visualizer.close()
                QMessageBox.information(self, "Успех", "График построен успешно!")

            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Ошибка при построении графика: {str(e)}")



if __name__ == '__main__':
    app = QApplication(sys.argv)
    ex = App()
    ex.show()
    sys.exit(app.exec())

